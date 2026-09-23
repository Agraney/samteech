import openpyxl

wb = openpyxl.load_workbook("MATERIAL PRODUCTION DATA WORKING 20.9.26.xlsm", data_only=True)

# Test PRODUCTION MASTER parsing
ws_pm = wb["PRODUCTION MASTER"]
month_cols = []
for c in range(3, 40, 3):
    m_name = ws_pm.cell(2, c).value
    if m_name:
        month_cols.append((c, str(m_name).strip()))

print("Parsed months:", [m[1] for m in month_cols])

pm_records = []
curr_sec = ""
for r in range(4, ws_pm.max_row + 1):
    c1 = ws_pm.cell(r, 1).value
    if c1:
        curr_sec = str(c1).strip()
    rating = ws_pm.cell(r, 2).value
    if not rating or not str(rating).strip().endswith("KVA"):
        continue
    rating = str(rating).strip()
    
    for c_idx, m_name in month_cols:
        target = ws_pm.cell(r, c_idx).value or 0
        actual = ws_pm.cell(r, c_idx + 1).value or 0
        variance = ws_pm.cell(r, c_idx + 2).value or 0
        
        # Clean numeric
        def to_f(v):
            if isinstance(v, (int, float)): return float(v)
            try: return float(str(v).replace(',', '').strip())
            except: return 0.0
            
        t_f, a_f, v_f = to_f(target), to_f(actual), to_f(variance)
        pm_records.append((curr_sec, rating, m_name, t_f, a_f, v_f))

print(f"Total PRODUCTION MASTER records extracted: {len(pm_records)}")
non_zero_pm = [p for p in pm_records if p[3] > 0 or p[4] > 0]
print(f"Non-zero records: {len(non_zero_pm)}")
for p in non_zero_pm:
    print(" ", p)

# Test CONSUMPTION parsing
ws_cons = wb["CONSUMPTION"]
cons_records = []
curr_rating = ""

for r in range(1, ws_cons.max_row + 1):
    c1 = str(ws_cons.cell(r, 1).value or '').strip()
    if c1.startswith("RATING:"):
        curr_rating = c1.replace("RATING:", "").strip()
        continue
    if not curr_rating or c1 in ("MATERIAL", ""):
        continue
        
    mat_name = c1
    m_type = str(ws_cons.cell(r, 2).value or '').strip()
    m_size = str(ws_cons.cell(r, 3).value or '').strip()
    pieces = ws_cons.cell(r, 4).value or 1
    qty_coil = ws_cons.cell(r, 5).value or 0
    qty_tr = ws_cons.cell(r, 6).value or 0
    unit = str(ws_cons.cell(r, 7).value or '').strip()
    
    def to_f(v):
        if isinstance(v, (int, float)): return float(v)
        try: return float(str(v).replace(',', '').strip())
        except: return 0.0
        
    cons_records.append((curr_rating, mat_name, m_type, m_size, to_f(pieces), to_f(qty_coil), to_f(qty_tr), unit))

print(f"\nTotal CONSUMPTION BOM records extracted: {len(cons_records)}")
for c in cons_records[:8]:
    print(" ", c)
