import openpyxl

wb = openpyxl.load_workbook("MATERIAL PRODUCTION DATA WORKING 20.9.26.xlsm", data_only=True)

# Inspect CONSUMPTION sheet in depth
ws_cons = wb["CONSUMPTION"]
print(f"--- CONSUMPTION Sheet (max_row={ws_cons.max_row}) ---")
for r in range(1, ws_cons.max_row + 1):
    vals = [ws_cons.cell(r, c).value for c in range(1, 9)]
    if any(vals):
        print(f"Row {r:2d}: {vals}")

# Inspect WIP Data in depth
ws_wip = wb["WIP Data"]
print(f"\n--- WIP Data Sheet (max_row={ws_wip.max_row}) ---")
for r in range(1, min(ws_wip.max_row + 1, 35)):
    vals = [ws_wip.cell(r, c).value for c in range(1, 15)]
    if any(vals):
        print(f"Row {r:2d}: {vals}")
