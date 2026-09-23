import openpyxl

wb = openpyxl.load_workbook("MATERIAL PRODUCTION DATA WORKING 20.9.26.xlsm", data_only=True)

print("All sheet names in workbook:")
for i, name in enumerate(wb.sheetnames, 1):
    print(f"{i:2d}. {name}")

# Detailed check of PRODUCTION MASTER
if "PRODUCTION MASTER" in wb.sheetnames:
    ws = wb["PRODUCTION MASTER"]
    print(f"\n================ PRODUCTION MASTER (max_row={ws.max_row}, max_col={ws.max_column}) ================")
    for r in range(1, 20):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 42))]
        if any(row_vals):
            print(f"Row {r:2d}: {[v for v in row_vals if v is not None][:12]}")

# Detailed check of CONSUMPTION
if "CONSUMPTION" in wb.sheetnames:
    ws = wb["CONSUMPTION"]
    print(f"\n================ CONSUMPTION (max_row={ws.max_row}, max_col={ws.max_column}) ================")
    for r in range(1, 25):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 20))]
        if any(row_vals):
            print(f"Row {r:2d}: {[v for v in row_vals if v is not None][:10]}")

# Detailed check of WIP Data
if "WIP Data" in wb.sheetnames:
    ws = wb["WIP Data"]
    print(f"\n================ WIP Data (max_row={ws.max_row}, max_col={ws.max_column}) ================")
    for r in range(1, 25):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 20))]
        if any(row_vals):
            print(f"Row {r:2d}: {[v for v in row_vals if v is not None][:10]}")

# Detailed check of Sheet1 and Sheet2
for s in ["Sheet1", "Sheet2"]:
    if s in wb.sheetnames:
        ws = wb[s]
        print(f"\n================ {s} (max_row={ws.max_row}, max_col={ws.max_column}) ================")
        for r in range(1, 10):
            row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 10))]
            if any(row_vals):
                print(f"Row {r:2d}: {[v for v in row_vals if v is not None]}")
