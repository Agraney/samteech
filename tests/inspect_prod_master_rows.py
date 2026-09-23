import openpyxl

wb = openpyxl.load_workbook("MATERIAL PRODUCTION DATA WORKING 20.9.26.xlsm", data_only=True)
ws_pm = wb["PRODUCTION MASTER"]

print("--- PRODUCTION MASTER Month Headers ---")
months = []
# Row 2 has Month names every 3 columns starting at col 3 (Apr-26, May-26, ...)
for c in range(3, 40, 3):
    m_name = ws_pm.cell(2, c).value
    if m_name:
        months.append((c, m_name))
print("Detected months:", months)

print("\n--- PRODUCTION MASTER Non-Zero Rows ---")
curr_section = ""
for r in range(4, ws_pm.max_row + 1):
    sec_val = ws_pm.cell(r, 1).value
    if sec_val:
        curr_section = sec_val.strip()
    rating = ws_pm.cell(r, 2).value
    if not rating:
        continue
        
    row_data = {}
    has_nonzero = False
    for col_idx, m_name in months:
        target = ws_pm.cell(r, col_idx).value or 0
        actual = ws_pm.cell(r, col_idx + 1).value or 0
        variance = ws_pm.cell(r, col_idx + 2).value or 0
        if actual != 0 or target != 0:
            has_nonzero = True
            row_data[m_name] = {"target": target, "actual": actual, "variance": variance}
            
    if has_nonzero:
        print(f"[{curr_section}] {rating}: {row_data}")
