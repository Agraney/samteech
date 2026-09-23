import hashlib
import datetime
import json
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl

from app.db import get_db, init_db
from app.config import UPLOADS_DIR

def compute_file_hash(file_path: Path) -> str:
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(8192):
            sha256.update(chunk)
    return sha256.hexdigest()

def clean_float(val: Any) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    # If string with commas or currency signs
    s = str(val).replace(",", "").replace("₹", "").replace("$", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0

def clean_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()

def clean_date(val: Any) -> Optional[str]:
    if val is None or val == "":
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    # Try various date formats
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d.%m.%Y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    # If string has space (datetime), take first part
    if " " in s:
        part = s.split(" ")[0]
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.datetime.strptime(part, fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
    return None


class WorkbookIngestor:
    def __init__(self, file_path: Path, notes: str = ""):
        self.file_path = file_path
        self.notes = notes
        self.wb = None

    def parse_and_store(self) -> Dict[str, Any]:
        init_db()
        file_hash = compute_file_hash(self.file_path)
        
        # Load workbook with data_only=True to read cached formula values
        self.wb = openpyxl.load_workbook(self.file_path, data_only=True, read_only=False)
        sheet_names = self.wb.sheetnames
        
        with get_db() as conn:
            # 1. Insert snapshot record
            cur = conn.execute(
                """
                INSERT INTO snapshots (filename, file_hash, file_path, notes)
                VALUES (?, ?, ?, ?)
                """,
                (self.file_path.name, file_hash, str(self.file_path), self.notes)
            )
            snapshot_id = cur.lastrowid
            
            # 2. Ingest BASIC sheet (Norms) if present
            norms_count = self._parse_basic_sheet(conn, snapshot_id)
            
            # 3. Ingest detailed CONSUMPTION BOM sheet if present
            bom_count = self._parse_consumption_sheet(conn, snapshot_id)
            
            # 4. Ingest PRODUCTION MASTER (Monthly targets & actuals) if present
            pm_count = self._parse_production_master(conn, snapshot_id)
            
            # 5. Ingest STOCK MASTER sheet
            materials_map, stock_master_count = self._parse_stock_master(conn, snapshot_id)
            
            # 6. Ingest STK_* transaction sheets
            transactions_count = self._parse_stk_sheets(conn, snapshot_id, materials_map)
            
            # 7. Ingest PROD_* daily production sheets
            prod_entries_count = self._parse_prod_sheets(conn, snapshot_id)
            
            summary = {
                "snapshot_id": snapshot_id,
                "filename": self.file_path.name,
                "file_hash": file_hash,
                "materials_count": len(materials_map),
                "stock_master_rows": stock_master_count,
                "transactions_count": transactions_count,
                "production_entries_count": prod_entries_count,
                "norms_count": norms_count,
                "bom_spec_count": bom_count,
                "monthly_targets_count": pm_count
            }
            
            # Save summary_json in snapshot
            conn.execute(
                "UPDATE snapshots SET summary_json = ? WHERE id = ?",
                (json.dumps(summary), snapshot_id)
            )
            
        return summary

    def _parse_production_master(self, conn, snapshot_id: int) -> int:
        if "PRODUCTION MASTER" not in self.wb.sheetnames:
            return 0
        ws_pm = self.wb["PRODUCTION MASTER"]
        
        month_cols = []
        for c in range(3, min(ws_pm.max_column + 1, 45), 3):
            m_name = ws_pm.cell(2, c).value
            if m_name:
                month_cols.append((c, clean_str(m_name)))
                
        count = 0
        curr_sec = ""
        for r in range(4, ws_pm.max_row + 1):
            c1 = ws_pm.cell(r, 1).value
            if c1:
                curr_sec = clean_str(c1)
                if "CORE COIL" in curr_sec:
                    curr_sec = "CORE COIL ASSEMBLY"
                elif "HV WINDING" in curr_sec:
                    curr_sec = "HV WINDING"
                elif "LV WINDING" in curr_sec:
                    curr_sec = "LV WINDING"
                elif "TANKING" in curr_sec:
                    curr_sec = "TANKING"
                elif "PASSED" in curr_sec:
                    curr_sec = "TESTING PASSED"
                elif "FAILED" in curr_sec:
                    curr_sec = "TESTING FAILED"
                elif "PAINTING" in curr_sec:
                    curr_sec = "PAINTING"
                elif "DISPATCH" in curr_sec:
                    curr_sec = "DISPATCH"
                    
            rating = ws_pm.cell(r, 2).value
            if not rating or not clean_str(rating).endswith("KVA"):
                continue
            rating_str = clean_str(rating)
            
            for c_idx, m_name in month_cols:
                target = clean_float(ws_pm.cell(r, c_idx).value)
                actual = clean_float(ws_pm.cell(r, c_idx + 1).value)
                variance = clean_float(ws_pm.cell(r, c_idx + 2).value)
                
                conn.execute(
                    """
                    INSERT INTO production_monthly_targets 
                    (snapshot_id, section, rating, month, target_units, actual_units, variance_units)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (snapshot_id, curr_sec, rating_str, m_name, target, actual, variance)
                )
                count += 1
        return count

    def _parse_consumption_sheet(self, conn, snapshot_id: int) -> int:
        if "CONSUMPTION" not in self.wb.sheetnames:
            return 0
        ws_cons = self.wb["CONSUMPTION"]
        count = 0
        curr_rating = ""
        
        for r in range(1, ws_cons.max_row + 1):
            c1 = clean_str(ws_cons.cell(r, 1).value)
            if c1.startswith("RATING:"):
                curr_rating = c1.replace("RATING:", "").strip()
                continue
            if not curr_rating or c1 in ("MATERIAL", ""):
                continue
                
            mat_name = c1
            m_type = clean_str(ws_cons.cell(r, 2).value)
            m_size = clean_str(ws_cons.cell(r, 3).value)
            pieces = clean_float(ws_cons.cell(r, 4).value) or 1.0
            qty_coil = clean_float(ws_cons.cell(r, 5).value)
            qty_tr = clean_float(ws_cons.cell(r, 6).value)
            unit = clean_str(ws_cons.cell(r, 7).value)
            
            conn.execute(
                """
                INSERT INTO bom_specifications 
                (snapshot_id, rating, material_name, material_type, size, pieces_count, qty_per_coil, qty_per_transformer, unit)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (snapshot_id, curr_rating, mat_name, m_type, m_size, pieces, qty_coil, qty_tr, unit)
            )
            count += 1
        return count


    def _parse_basic_sheet(self, conn, snapshot_id: int) -> int:
        if "BASIC" not in self.wb.sheetnames:
            return 0
        ws = self.wb["BASIC"]
        
        # Find header row
        header_row = None
        col_map = {}
        for r in range(1, 10):
            row_vals = [clean_str(ws.cell(r, c).value).upper() for c in range(1, 10)]
            if any("RATING" in v for v in row_vals) and any("QTY" in v or "NORM" in v for v in row_vals):
                header_row = r
                for c in range(1, 10):
                    val = clean_str(ws.cell(r, c).value).upper()
                    if "RATING" in val:
                        col_map["rating"] = c
                    elif "MATERIAL" in val:
                        col_map["material"] = c
                    elif "QTY" in val or "NORM" in val:
                        col_map["qty"] = c
                    elif "UNIT" in val:
                        col_map["unit"] = c
                break
                
        if not header_row or "rating" not in col_map or "qty" not in col_map:
            return 0
            
        count = 0
        for r in range(header_row + 1, ws.max_row + 1):
            rating = clean_str(ws.cell(r, col_map["rating"]).value)
            mat_name = clean_str(ws.cell(r, col_map.get("material", 2)).value)
            qty = clean_float(ws.cell(r, col_map["qty"]).value)
            unit = clean_str(ws.cell(r, col_map.get("unit", 4)).value)
            
            if rating and mat_name and qty > 0:
                conn.execute(
                    """
                    INSERT INTO consumption_norms (snapshot_id, rating, material_name, qty_per_transformer, unit)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (snapshot_id, rating, mat_name, qty, unit)
                )
                count += 1
        return count

    def _parse_stock_master(self, conn, snapshot_id: int) -> tuple[Dict[str, int], int]:
        """
        Parses STOCK MASTER.
        Columns: Rating (B=2), Sheet name (C=3), Material (D=4), Type (E=5), Size (F=6),
        Unit (G=7), Opening Balance (H=8), Received Qty (I=9), Rate (J=10), Value (K=11),
        Issued Qty (O=15), Closing Balance (Q=17)
        """
        if "STOCK MASTER" not in self.wb.sheetnames:
            return {}, 0
            
        ws = self.wb["STOCK MASTER"]
        materials_map = {} # (rating, sheet_name) -> material_id
        count = 0
        
        # Scan header
        header_row = None
        for r in range(1, 10):
            val_b = clean_str(ws.cell(r, 2).value).upper()
            val_c = clean_str(ws.cell(r, 3).value).upper()
            if "RATING" in val_b or "SHEET" in val_c:
                header_row = r
                break
        start_row = (header_row + 1) if header_row else 5
        
        for r in range(start_row, ws.max_row + 1):
            rating = clean_str(ws.cell(r, 2).value)
            sheet_name = clean_str(ws.cell(r, 3).value)
            mat_name = clean_str(ws.cell(r, 4).value)
            
            if not rating or not sheet_name or not mat_name:
                continue
                
            mat_type = clean_str(ws.cell(r, 5).value)
            size = clean_str(ws.cell(r, 6).value)
            unit = clean_str(ws.cell(r, 7).value)
            op_bal = clean_float(ws.cell(r, 8).value)
            rec_qty = clean_float(ws.cell(r, 9).value)
            rate = clean_float(ws.cell(r, 10).value)
            val = clean_float(ws.cell(r, 11).value)
            iss_qty = clean_float(ws.cell(r, 15).value)
            cl_bal = clean_float(ws.cell(r, 17).value)
            
            # Upsert into materials catalog
            cur = conn.execute(
                """
                INSERT INTO materials (rating, sheet_name, material_name, material_type, size, unit)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(rating, sheet_name, material_name) DO UPDATE SET
                    material_type = excluded.material_type,
                    size = excluded.size,
                    unit = excluded.unit
                RETURNING id;
                """,
                (rating, sheet_name, mat_name, mat_type, size, unit)
            )
            row = cur.fetchone()
            if row:
                material_id = row[0]
            else:
                cur2 = conn.execute(
                    "SELECT id FROM materials WHERE rating = ? AND sheet_name = ? AND material_name = ?",
                    (rating, sheet_name, mat_name)
                )
                material_id = cur2.fetchone()[0]
                
            materials_map[(rating, sheet_name)] = material_id
            materials_map[sheet_name] = material_id # convenience lookup
            
            # Record in stock_master_snapshots
            conn.execute(
                """
                INSERT INTO stock_master_snapshots 
                (snapshot_id, material_id, opening_balance, received_qty, rate, value, issued_qty, closing_balance)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (snapshot_id, material_id, op_bal, rec_qty, rate, val, iss_qty, cl_bal)
            )
            count += 1
            
        return materials_map, count

    def _parse_stk_sheets(self, conn, snapshot_id: int, materials_map: Dict[Any, int]) -> int:
        """
        Parses ~72 STK_* sheets.
        Header on row 5:
        S.NO (1) | DATE (2) | OPENING BALANCE (3) | RECEIVED QTY (4) | RATE (₹) (5) | VALUE (₹) (6) |
        SUPPLIER NAME (7) | INVOICE NO (8) | ISSUED QTY (9) | ISSUED TO SECTION (10) | CLOSING BALANCE (11) | REMARKS (12)
        """
        total_tx = 0
        for sheet_name in self.wb.sheetnames:
            if not sheet_name.startswith("STK_"):
                continue
                
            material_id = materials_map.get(sheet_name)
            if not material_id:
                # If not mapped from STOCK MASTER, auto-register material
                parts = sheet_name.split("_")
                rating = parts[1] if len(parts) > 1 and parts[1].endswith("KVA") else "COMMON"
                mat_name = "_".join(parts[2:]) if rating != "COMMON" and len(parts) > 2 else "_".join(parts[1:])
                cur = conn.execute(
                    """
                    INSERT INTO materials (rating, sheet_name, material_name, material_type, size, unit)
                    VALUES (?, ?, ?, 'Generic', '', 'UNIT')
                    ON CONFLICT(rating, sheet_name, material_name) DO UPDATE SET id=id
                    RETURNING id;
                    """,
                    (rating, sheet_name, mat_name)
                )
                row = cur.fetchone()
                material_id = row[0] if row else conn.execute(
                    "SELECT id FROM materials WHERE rating = ? AND sheet_name = ?", (rating, sheet_name)
                ).fetchone()[0]
                materials_map[sheet_name] = material_id
                
            ws = self.wb[sheet_name]
            
            # Find header row
            header_row = 5
            for r in range(1, 8):
                v1 = clean_str(ws.cell(r, 1).value).upper()
                v2 = clean_str(ws.cell(r, 2).value).upper()
                if "S.NO" in v1 or "DATE" in v2:
                    header_row = r
                    break
                    
            for r in range(header_row + 1, ws.max_row + 1):
                s_no_val = ws.cell(r, 1).value
                date_val = clean_date(ws.cell(r, 2).value)
                op_bal = clean_float(ws.cell(r, 3).value)
                rec_qty = clean_float(ws.cell(r, 4).value)
                rate = clean_float(ws.cell(r, 5).value)
                val = clean_float(ws.cell(r, 6).value)
                supplier = clean_str(ws.cell(r, 7).value)
                invoice = clean_str(ws.cell(r, 8).value)
                iss_qty = clean_float(ws.cell(r, 9).value)
                section = clean_str(ws.cell(r, 10).value)
                cl_bal = clean_float(ws.cell(r, 11).value)
                remarks = clean_str(ws.cell(r, 12).value)
                
                # Only record rows that have actual transactions or movement
                if rec_qty == 0 and iss_qty == 0 and not supplier and not invoice and not remarks:
                    continue
                    
                s_no = int(s_no_val) if isinstance(s_no_val, (int, float)) and s_no_val > 0 else None
                
                conn.execute(
                    """
                    INSERT INTO stock_transactions (
                        snapshot_id, material_id, s_no, date, opening_balance,
                        received_qty, rate, value, supplier_name, invoice_no,
                        issued_qty, issued_to_section, closing_balance, remarks
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        snapshot_id, material_id, s_no, date_val, op_bal,
                        rec_qty, rate, val, supplier, invoice,
                        iss_qty, section, cl_bal, remarks
                    )
                )
                total_tx += 1
                
        return total_tx

    def _parse_prod_sheets(self, conn, snapshot_id: int) -> int:
        """
        Parses 8 PROD_<SECTION> sheets.
        Date on Col A (or B), ratings on columns, values = units produced.
        Terminates parsing before the MONTHLY SUMMARY table to prevent double counting.
        """
        total_prod = 0
        ratings_standard = ["16KVA", "25KVA", "63KVA", "100KVA", "250KVA"]
        
        for sheet_name in self.wb.sheetnames:
            if not sheet_name.startswith("PROD_") or sheet_name in ("PROD_MASTER", "PRODUCTION MASTER"):
                continue
                
            section_name = sheet_name.replace("PROD_", "").strip()
            # Normalize naming
            if section_name in ("CORE COIL ASSY", "CORE COIL ASSEMBLY"):
                section_name = "CORE COIL ASSEMBLY"
                
            ws = self.wb[sheet_name]
            
            # Find header row
            header_row = 2
            col_date = 1
            rating_cols = {} # rating -> col_idx
            
            for r in range(1, 8):
                row_vals = [clean_str(ws.cell(r, c).value).upper() for c in range(1, 15)]
                if any("DATE" in v for v in row_vals):
                    header_row = r
                    for c in range(1, 15):
                        v = clean_str(ws.cell(r, c).value).upper()
                        if "DATE" in v:
                            col_date = c
                        for rat in ratings_standard:
                            if rat in v and rat not in rating_cols:
                                rating_cols[rat] = c
                    break
                    
            if not rating_cols:
                # Default mapping if headers are standard
                rating_cols = {
                    "16KVA": 2, "25KVA": 3, "63KVA": 4, "100KVA": 5, "250KVA": 6
                }
                
            for r in range(header_row + 1, ws.max_row + 1):
                # Check for MONTHLY SUMMARY or TOTAL row - terminate daily log parsing
                c1_str = clean_str(ws.cell(r, col_date).value).upper()
                if "SUMMARY" in c1_str or c1_str == "MONTH" or c1_str.startswith("TOTAL"):
                    break
                    
                raw_date = ws.cell(r, col_date).value
                date_str = clean_date(raw_date)
                if not date_str:
                    continue
                    
                has_any = False
                for rating, c_idx in rating_cols.items():
                    units = clean_float(ws.cell(r, c_idx).value)
                    if units > 0:
                        conn.execute(
                            """
                            INSERT INTO production_entries (
                                snapshot_id, date, section, rating, units_produced
                            ) VALUES (?, ?, ?, ?, ?)
                            """,
                            (snapshot_id, date_str, section_name, rating, units)
                        )
                        total_prod += 1
                        has_any = True
                        
        return total_prod

